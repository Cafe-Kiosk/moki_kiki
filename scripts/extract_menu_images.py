"""엑셀 사진메뉴판에 포함된 33개 이미지를 웹 자산으로 추출한다."""
from __future__ import annotations

import posixpath
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "data" / "moki_data.xlsx"
OUTPUT = ROOT / "assets" / "menu"

NS = {
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
}


def extract() -> list[Path]:
    """사진메뉴판 배치 순서를 읽어 메뉴 ID 이름으로 JPEG를 저장한다."""
    workbook = load_workbook(SOURCE, data_only=True)
    sheet = workbook["사진메뉴판"]
    OUTPUT.mkdir(parents=True, exist_ok=True)

    with zipfile.ZipFile(SOURCE) as archive:
        relationships = ET.fromstring(
            archive.read("xl/drawings/_rels/drawing1.xml.rels")
        )
        targets = {
            relation.attrib["Id"]: relation.attrib["Target"]
            for relation in relationships.findall("rel:Relationship", NS)
        }
        drawing = ET.fromstring(archive.read("xl/drawings/drawing1.xml"))
        written: list[Path] = []

        for anchor in drawing:
            origin = anchor.find("xdr:from", NS)
            blip = anchor.find(".//a:blip", NS)
            if origin is None or blip is None:
                continue

            row = int(origin.findtext("xdr:row", namespaces=NS) or 0) + 1
            column = int(origin.findtext("xdr:col", namespaces=NS) or 0) + 1
            menu_id = str(sheet.cell(row=row + 15, column=column + 2).value or "")
            if not re.fullmatch(r"M\d{2}", menu_id):
                raise ValueError(f"사진 위치 {row},{column}에서 메뉴 ID를 찾지 못했습니다")

            relation_id = blip.attrib[f"{{{NS['r']}}}embed"]
            member = posixpath.normpath(
                posixpath.join("xl/drawings", targets[relation_id])
            )
            destination = OUTPUT / f"{menu_id}.jpg"
            destination.write_bytes(archive.read(member))
            written.append(destination)

    if len(written) != 33 or len({path.stem for path in written}) != 33:
        raise ValueError(f"메뉴 사진은 33장이어야 합니다: {len(written)}장")
    return written


if __name__ == "__main__":
    files = extract()
    print(f"메뉴 사진 {len(files)}장을 {OUTPUT.relative_to(ROOT)}에 추출했습니다.")
